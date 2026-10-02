"""Generates One_Levantamento_e_Requisitos.docx and .pdf from conteudo.md, metadados.json and diagramas/.

Usage: python docs/requisitos/gerar.py
Needs python-docx (requirements.txt), Java 17 for PlantUML, and Microsoft Word to update the
table of contents and export the PDF.
"""

import json
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import urllib.error
import urllib.request
from pathlib import Path

import docx
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
NOME = "One_Levantamento_e_Requisitos"
VERSAO_PLANTUML = "1.2026.8"
JAR = AQUI / ".cache" / f"plantuml-{VERSAO_PLANTUML}.jar"
URL_JAR = (
    f"https://github.com/plantuml/plantuml/releases/download/v{VERSAO_PLANTUML}/plantuml-{VERSAO_PLANTUML}.jar"
)
TIMES = "Times New Roman"
LARGURA_TEXTO = Cm(16)
ALTURA_MAX_FIGURA = Cm(21)
LINHAS_QUADRO_INTEIRO = 12
LARGURA_LOGO = Cm(5)
FONTE_FIGURA = "Fonte: elaborado pelos autores (2026)."

SCRIPT_WORD = r"""
param([string]$Docx, [string]$Pdf)
$ErrorActionPreference = "Stop"
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
try {
    $doc = $word.Documents.Open($Docx)
    $doc.Fields.Update() | Out-Null
    foreach ($sumario in $doc.TablesOfContents) { $sumario.Update() }
    $doc.Save()
    $doc.ExportAsFixedFormat($Pdf, 17)
    $doc.Close()
}
finally {
    $word.Quit()
}
"""


def java_disponivel():
    return shutil.which("java") is not None


def word_disponivel():
    try:
        resultado = subprocess.run(
            ["powershell", "-NoProfile", "-Command",
             "try { (New-Object -ComObject Word.Application).Quit(); exit 0 } catch { exit 1 }"],
            capture_output=True, timeout=120,
        )
    except (OSError, subprocess.TimeoutExpired):
        return False
    return resultado.returncode == 0


def obter_plantuml():
    if not JAR.is_file():
        JAR.parent.mkdir(exist_ok=True)
        parcial = JAR.with_suffix(".parcial")
        urllib.request.urlretrieve(URL_JAR, parcial)
        parcial.replace(JAR)
    return JAR


def renderizar(jar, destino):
    fontes = sorted((AQUI / "diagramas").glob("*.puml"))
    subprocess.run(
        ["java", "-Djava.awt.headless=true", "-jar", str(jar), "-charset", "UTF-8", "-tpng",
         "-o", str(destino), *map(str, fontes)],
        check=True,
    )
    return {f.stem: destino / f"{f.stem}.png" for f in fontes}


# --- conteudo.md ---------------------------------------------------------------------------

def ler_conteudo(texto):
    """Returns (resumo, palavras_chave, blocos) from the Markdown subset conteudo.md uses."""
    linhas = texto.splitlines()
    resumo, palavras, blocos, i = [], "", [], 0

    def tabela(inicio):
        fim = inicio
        while fim < len(linhas) and linhas[fim].startswith("|"):
            fim += 1
        celulas = [[c.strip() for c in l.strip().strip("|").split("|")] for l in linhas[inicio:fim]]
        return [c for c in celulas if not all(re.fullmatch(r":?-+:?", x) for x in c)], fim

    while i < len(linhas):
        linha = linhas[i].strip()
        if not linha:
            i += 1
        elif linha.startswith("#"):
            nivel = len(linha) - len(linha.lstrip("#"))
            blocos.append(("titulo", nivel, linha[nivel:].strip()))
            i += 1
        elif linha.startswith("Quadro:"):
            linhas_tabela, i = tabela(i + 1)
            blocos.append(("quadro", linha[len("Quadro:"):].strip(), linhas_tabela, linhas[i].strip()))
            i += 1
        elif linha.startswith("Figura:"):
            imagem = re.match(r"!\[[^\]]*\]\(diagramas/([^)]+)\.png\)", linhas[i + 1].strip()).group(1)
            blocos.append(("figura", linha[len("Figura:"):].strip(), imagem, linhas[i + 2].strip()))
            i += 3
        elif linha.startswith("- "):
            itens = []
            while i < len(linhas) and linhas[i].startswith("- "):
                itens.append(linhas[i][2:].strip())
                i += 1
            blocos.append(("lista", itens))
        elif linha.startswith("Palavras-chave:"):
            palavras = linha
            i += 1
        else:
            paragrafo = []
            while i < len(linhas) and linhas[i].strip() and not linhas[i].startswith(("#", "- ", "|")):
                paragrafo.append(linhas[i].strip())
                i += 1
            if blocos:
                blocos.append(("paragrafo", " ".join(paragrafo)))
            else:
                resumo.append(" ".join(paragrafo))
    return resumo, palavras, blocos


# --- estilos e página ----------------------------------------------------------------------

def times_em(rfonts):
    for atributo in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(atributo), None)
    for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(atributo), TIMES)


def formatar(estilo, tamanho=None, negrito=None, alinhamento=None, recuo=None, linha=None,
             antes=None, depois=None):
    times_em(estilo.element.get_or_add_rPr().get_or_add_rFonts())
    estilo.font.color.rgb = RGBColor(0, 0, 0)
    estilo.font.italic = False
    if tamanho:
        estilo.font.size = Pt(tamanho)
    if negrito is not None:
        estilo.font.bold = negrito
    formato = estilo.paragraph_format
    for atributo, valor in (("alignment", alinhamento), ("first_line_indent", recuo), ("line_spacing", linha),
                            ("space_before", antes), ("space_after", depois)):
        if valor is not None:
            setattr(formato, atributo, valor)
    return estilo


def configurar_estilos(doc):
    padrao = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    times_em(padrao.find(qn("w:rFonts")))
    idioma = padrao.find(qn("w:lang")) if padrao.find(qn("w:lang")) is not None else OxmlElement("w:lang")
    idioma.set(qn("w:val"), "pt-BR")
    padrao.append(idioma)

    centro, esquerda = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT
    formatar(doc.styles["Normal"], 12, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, recuo=Cm(1.25), linha=1.5,
             antes=Pt(0), depois=Pt(12))
    for nivel, negrito in ((1, True), (2, False), (3, True)):
        estilo = formatar(doc.styles[f"Heading {nivel}"], 12, negrito, esquerda, Cm(0), 1.5, Pt(12), Pt(12))
        estilo.paragraph_format.keep_with_next = True
        estilo.paragraph_format.page_break_before = nivel == 1
    lista = formatar(doc.styles["List Bullet"], depois=Pt(6))
    lista.paragraph_format.left_indent = Cm(1.25)
    lista.paragraph_format.first_line_indent = Cm(-0.63)

    novos = {
        "Capa": dict(alinhamento=centro, recuo=Cm(0), depois=Pt(0)),
        "Natureza": dict(alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, recuo=Cm(0), linha=1.0, depois=Pt(12)),
        "Titulo sem numero": dict(negrito=True, alinhamento=centro, recuo=Cm(0), depois=Pt(12)),
        "Legenda": dict(tamanho=10, alinhamento=esquerda, recuo=Cm(0), linha=1.0, antes=Pt(6), depois=Pt(3)),
        "Fonte": dict(tamanho=10, alinhamento=esquerda, recuo=Cm(0), linha=1.0, antes=Pt(3), depois=Pt(12)),
        "Tabela": dict(tamanho=10, alinhamento=esquerda, recuo=Cm(0), linha=1.0, depois=Pt(0)),
        "Figura": dict(alinhamento=centro, recuo=Cm(0), linha=1.0, depois=Pt(0)),
        "Referencia": dict(alinhamento=esquerda, recuo=Cm(0), linha=1.0, depois=Pt(12)),
    }
    for nome, formato in novos.items():
        estilo = doc.styles.add_style(nome, WD_STYLE_TYPE.PARAGRAPH)
        estilo.base_style = doc.styles["Normal"]
        formatar(estilo, **formato)
    doc.styles["Natureza"].paragraph_format.left_indent = Cm(8)
    doc.styles["Titulo sem numero"].paragraph_format.page_break_before = True
    doc.styles["Legenda"].paragraph_format.keep_with_next = True
    doc.styles["Figura"].paragraph_format.keep_with_next = True


def pagina_a4(secao):
    secao.page_width, secao.page_height = Cm(21), Cm(29.7)
    secao.top_margin = secao.left_margin = Cm(3)
    secao.bottom_margin = secao.right_margin = Cm(2)
    secao.header_distance = Cm(2)


def reiniciar_numeracao(secao):
    numeracao = OxmlElement("w:pgNumType")
    numeracao.set(qn("w:start"), "1")
    # invariant: CT_SectPr is an ordered sequence; pgNumType must precede cols/docGrid or Word
    # reports the file as damaged.
    seguinte = next((filho for filho in secao._sectPr if filho.tag in (qn("w:cols"), qn("w:docGrid"))), None)
    if seguinte is not None:
        seguinte.addprevious(numeracao)
    else:
        secao._sectPr.append(numeracao)


def campo(paragrafo, instrucao, provisorio):
    def marcador(tipo):
        elemento = OxmlElement("w:fldChar")
        elemento.set(qn("w:fldCharType"), tipo)
        paragrafo.add_run()._r.append(elemento)

    marcador("begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instrucao
    paragrafo.add_run()._r.append(instr)
    marcador("separate")
    paragrafo.add_run(provisorio)
    marcador("end")


# --- partes do documento -------------------------------------------------------------------

def texto_com_negrito(paragrafo, texto):
    for n, trecho in enumerate(texto.split("**")):
        if trecho:
            paragrafo.add_run(trecho).bold = n % 2 == 1


def linha_central(doc, texto, negrito=False, antes=0):
    paragrafo = doc.add_paragraph(style="Capa")
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragrafo.paragraph_format.space_before = Pt(antes)
    paragrafo.add_run(texto).bold = negrito
    return paragrafo


def capa_e_folha_de_rosto(doc, meta):
    logo = doc.add_paragraph(style="Capa")
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo.add_run().add_picture(str(AQUI / meta["logo"]), width=LARGURA_LOGO)
    linha_central(doc, meta["instituicao"], negrito=True, antes=12)
    linha_central(doc, meta["curso"], negrito=True)
    for n, autor in enumerate(meta["autores"]):
        linha_central(doc, autor, antes=96 if n == 0 else 0)
    linha_central(doc, meta["titulo"], negrito=True, antes=120)
    linha_central(doc, meta["subtitulo"])
    linha_central(doc, meta["cidade"], antes=150)
    linha_central(doc, meta["ano"])

    secao = doc.add_section(WD_SECTION.NEW_PAGE)
    pagina_a4(secao)
    reiniciar_numeracao(secao)
    for autor in meta["autores"]:
        linha_central(doc, autor)
    linha_central(doc, meta["titulo"], negrito=True, antes=120)
    linha_central(doc, meta["subtitulo"])
    natureza = doc.add_paragraph(meta["natureza"], style="Natureza")
    natureza.paragraph_format.space_before = Pt(60)
    doc.add_paragraph(meta["orientador"], style="Natureza")
    linha_central(doc, meta["cidade"], antes=150)
    linha_central(doc, meta["ano"])


def resumo_e_sumario(doc, resumo, palavras):
    doc.add_paragraph("RESUMO", style="Titulo sem numero")
    for texto in resumo:
        doc.add_paragraph(texto)
    chave = doc.add_paragraph()
    rotulo, termos = palavras.split(":", 1)
    chave.add_run(rotulo + ":").bold = True
    chave.add_run(termos)
    doc.add_paragraph("SUMÁRIO", style="Titulo sem numero")
    campo(doc.add_paragraph(style="Normal"), 'TOC \\o "1-3" \\h \\z \\u', "Atualize o sumário no Word.")


def larguras(linhas):
    """Each column fits its longest word; the rest of the width follows the text length."""
    colunas = range(len(linhas[0]))
    minimos = [Cm(0.22 * max(len(p) for linha in linhas for p in re.split(r"\s+|<br>", linha[c])) + 0.4)
               for c in colunas]
    if sum(minimos) > LARGURA_TEXTO:
        return [int(LARGURA_TEXTO * m / sum(minimos)) for m in minimos]
    pesos = [max(min(len(linha[c]), 60) for linha in linhas) for c in colunas]
    sobra = LARGURA_TEXTO - sum(minimos)
    return [int(m + sobra * p / sum(pesos)) for m, p in zip(minimos, pesos)]


def quadro(doc, numero, titulo, linhas, fonte):
    doc.add_paragraph(f"Quadro {numero} – {titulo}", style="Legenda")
    tabela = doc.add_table(rows=len(linhas), cols=len(linhas[0]))
    tabela.style = doc.styles["Table Grid"]
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.autofit = False
    medidas = larguras(linhas)
    for r, linha in enumerate(linhas):
        for c, texto in enumerate(linha):
            celula = tabela.cell(r, c)
            celula.width = medidas[c]
            partes = texto.split("<br>")
            celula.paragraphs[0].style = doc.styles["Tabela"]
            for n, parte in enumerate(partes):
                paragrafo = celula.paragraphs[0] if n == 0 else celula.add_paragraph(style="Tabela")
                paragrafo.add_run(parte).bold = r == 0
                # why: a quadro short enough for one page stays whole, with its caption and source.
                paragrafo.paragraph_format.keep_with_next = len(linhas) <= LINHAS_QUADRO_INTEIRO
    cabecalho = tabela.rows[0]._tr.get_or_add_trPr()
    repetir = OxmlElement("w:tblHeader")
    cabecalho.append(repetir)
    doc.add_paragraph(fonte, style="Fonte")


def figura(doc, numero, titulo, imagem, fonte):
    legenda = doc.add_paragraph(f"Figura {numero} – {titulo}", style="Legenda")
    legenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
    largura_px, altura_px = struct.unpack(">II", imagem.read_bytes()[16:24])
    largura = min(LARGURA_TEXTO, int(ALTURA_MAX_FIGURA * largura_px / altura_px))
    doc.add_paragraph(style="Figura").add_run().add_picture(str(imagem), width=largura)
    rodape = doc.add_paragraph(fonte, style="Fonte")
    rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER


def corpo(doc, blocos, imagens):
    secao = doc.add_section(WD_SECTION.NEW_PAGE)
    pagina_a4(secao)
    numeracao = secao._sectPr.find(qn("w:pgNumType"))
    if numeracao is not None:
        secao._sectPr.remove(numeracao)
    secao.header.is_linked_to_previous = False
    pagina = secao.header.paragraphs[0]
    pagina.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pagina.paragraph_format.first_line_indent = Cm(0)
    campo(pagina, "PAGE", "1")

    contadores, quadros, figuras, em_referencias = [0, 0, 0], 0, 0, False
    for bloco in blocos:
        tipo = bloco[0]
        if tipo == "titulo":
            nivel, texto = bloco[1], bloco[2]
            if texto.casefold() == "referências":
                em_referencias = True
                doc.add_paragraph(texto.upper(), style="Heading 1").alignment = WD_ALIGN_PARAGRAPH.CENTER
                continue
            contadores[nivel - 1] += 1
            contadores[nivel:] = [0] * (3 - nivel)
            numero = ".".join(str(c) for c in contadores[:nivel])
            doc.add_paragraph(f"{numero} {texto.upper() if nivel <= 2 else texto}", style=f"Heading {nivel}")
        elif tipo == "paragrafo":
            texto_com_negrito(doc.add_paragraph(style="Referencia" if em_referencias else "Normal"), bloco[1])
        elif tipo == "lista":
            for item in bloco[1]:
                texto_com_negrito(doc.add_paragraph(style="List Bullet"), item)
        elif tipo == "quadro":
            quadros += 1
            quadro(doc, quadros, bloco[1], bloco[2], bloco[3])
        elif tipo == "figura":
            figuras += 1
            figura(doc, figuras, bloco[1], imagens[bloco[2]], bloco[3])


def montar_documento(destino, imagens):
    meta = json.loads((AQUI / "metadados.json").read_text(encoding="utf-8"))
    resumo, palavras, blocos = ler_conteudo((AQUI / "conteudo.md").read_text(encoding="utf-8"))
    doc = docx.Document()
    doc.core_properties.title = f"{meta['titulo']}: {meta['subtitulo']}"
    doc.core_properties.author = "; ".join(meta["autores"])
    configurar_estilos(doc)
    pagina_a4(doc.sections[0])
    capa_e_folha_de_rosto(doc, meta)
    resumo_e_sumario(doc, resumo, palavras)
    corpo(doc, blocos, imagens)
    doc.save(str(destino))


def finalizar_no_word(docx_tmp, pdf_tmp):
    script = docx_tmp.with_name("word.ps1")
    script.write_text(SCRIPT_WORD, encoding="utf-8")
    subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script),
         "-Docx", str(docx_tmp), "-Pdf", str(pdf_tmp)],
        check=True,
    )


def main():
    faltando = [nome for nome, disponivel in (("Java", java_disponivel), ("Microsoft Word", word_disponivel))
                if not disponivel()]
    if faltando:
        print("Ferramenta ausente: " + ", ".join(faltando) + ". Nada foi alterado.", file=sys.stderr)
        return 1
    try:
        with tempfile.TemporaryDirectory() as temporario:
            temporario = Path(temporario)
            imagens = renderizar(obter_plantuml(), temporario)
            docx_tmp, pdf_tmp = temporario / f"{NOME}.docx", temporario / f"{NOME}.pdf"
            montar_documento(docx_tmp, imagens)
            finalizar_no_word(docx_tmp, pdf_tmp)
            # why: outputs are replaced only after every step succeeded, so a failure leaves the
            # previous docx/pdf in place.
            for imagem in imagens.values():
                shutil.copy2(imagem, AQUI / "diagramas" / imagem.name)
            shutil.copy2(docx_tmp, AQUI / f"{NOME}.docx")
            shutil.copy2(pdf_tmp, AQUI / f"{NOME}.pdf")
    except (OSError, subprocess.CalledProcessError, urllib.error.URLError) as erro:
        print(f"Falha ao gerar o documento: {erro}. Os arquivos anteriores foram mantidos.", file=sys.stderr)
        return 1
    print(AQUI / f"{NOME}.docx")
    print(AQUI / f"{NOME}.pdf")
    return 0


if __name__ == "__main__":
    sys.exit(main())
